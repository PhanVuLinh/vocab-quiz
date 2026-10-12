"""
Extract vocabulary data from PPTX files → vocab.json + images/
Unified parser: handles all PPTX formats by analyzing text content intelligently.
"""

import os
import sys
import io
import json
import re
from pptx import Presentation

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PPTX_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'pp')
OUTPUT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(OUTPUT_DIR, 'data')
IMAGES_DIR = os.path.join(OUTPUT_DIR, 'images')

# IPA pattern
IPA_PATTERN = re.compile(r'/[^/]{2,}/')
# Vietnamese chars
VN_CHARS = set('àáạảãâầấậẩẫăằắặẳẵèéẹẻẽêềếệểễìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũưừứựửữỳýỵỷỹđÀÁẠẢÃÂẦẤẬẨẪĂẰẮẶẲẴÈÉẸẺẼÊỀẾỆỂỄÌÍỊỈĨÒÓỌỎÕÔỒỐỘỔỖƠỜỚỢỞỠÙÚỤỦŨƯỪỨỰỬỮỲÝỴỶỸĐ')
# Title keywords to skip
TITLE_KEYWORDS = ['VOCABULARY', 'TỪ VỰNG', 'LESSON', 'LISTENING', 'READING', 'ĐỌC', 'NGHE', 'Phần', 'LUYỆN TẬP', 'PART']


def ensure_dirs():
    os.makedirs(DATA_DIR, exist_ok=True)
    os.makedirs(IMAGES_DIR, exist_ok=True)


def extract_lesson_number(filename):
    match = re.search(r'(\d+)', filename)
    return int(match.group(1)) if match else 0


def has_vietnamese(text):
    return bool(set(text) & VN_CHARS)


def is_title_slide(texts):
    """Check if this slide is a title/header slide (no vocab content)."""
    if not texts:
        return True
    combined = ' '.join(texts).upper().strip()
    # Pure title slides: only contain lesson name/title keywords
    for kw in TITLE_KEYWORDS:
        combined = combined.replace(kw.upper(), '').strip()
    # Remove digits, whitespace, punctuation
    cleaned = re.sub(r'[\d\s,&\-–]+', '', combined)
    # Remove "PART X" pattern
    cleaned = re.sub(r'PART\s*\d*', '', cleaned, flags=re.IGNORECASE)
    return len(cleaned) < 3  # Nearly empty after removing keywords = title slide


def save_image(shape, lesson_id, word_index):
    try:
        image = shape.image
        ext = image.content_type.split('/')[-1]
        if ext == 'jpeg':
            ext = 'jpg'
        filename = "b%02d_w%02d.%s" % (lesson_id, word_index, ext)
        filepath = os.path.join(IMAGES_DIR, filename)
        with open(filepath, 'wb') as f:
            f.write(image.blob)
        return "images/%s" % filename
    except Exception:
        return None


def extract_ipa(text):
    """Extract all IPA pronunciations from text."""
    matches = IPA_PATTERN.findall(text)
    return ' '.join(matches) if matches else ''


def split_example(text):
    """Split example into English and Vietnamese."""
    text = text.strip()
    # Find Vietnamese translation in parentheses at end
    paren_match = re.search(r'\(([^)]*[' + ''.join(list(VN_CHARS)[:10]) + r'][^)]*)\)\s*\.?\s*$', text)
    if paren_match:
        vi = paren_match.group(1).strip()
        en = text[:paren_match.start()].strip().rstrip('.')
        return en, vi
    return text, ''


def parse_slide_unified(slide, lesson_id, word_index):
    """
    Universal slide parser. Collects all text and images from a slide,
    then intelligently extracts: word, part of speech, IPA, meaning, example.
    """
    text_blocks = []
    images = []
    
    for shape in slide.shapes:
        if shape.shape_type == 13:  # Picture
            images.append(shape)
        elif hasattr(shape, 'text'):
            raw = shape.text.strip()
            if raw:
                # Replace vertical tab with newline
                raw = raw.replace('\x0b', '\n')
                # Skip slide number placeholders
                cleaned = raw.replace('‹#›', '').strip()
                if cleaned and not re.match(r'^\d{1,3}$', cleaned):
                    text_blocks.append(cleaned)
    
    if not text_blocks:
        return None
    
    # Check if this is a title slide
    if is_title_slide(text_blocks):
        return None
    
    # Strategy: combine all text blocks, then parse lines
    # Some slides have everything in 1 block, some in separate blocks
    all_lines = []
    for block in text_blocks:
        lines = [l.strip() for l in block.split('\n') if l.strip()]
        all_lines.extend(lines)
    
    if not all_lines:
        return None
    
    word = ''
    pos = ''
    pronunciation = ''
    meaning = ''
    example_en = ''
    example_vi = ''
    
    # ===== PASS 1: Find the word line (contains IPA or numbered entry) =====
    word_line_idx = -1
    
    for i, line in enumerate(all_lines):
        # Check for numbered word: "1. cash (n)" or "1. consider (v) /kənˈsɪdər/"
        num_match = re.match(r'^(\d+)\.\s*(.+)', line)
        if num_match:
            word_line_idx = i
            word_part = num_match.group(2).strip()
            # Extract IPA if on same line
            if IPA_PATTERN.search(word_part):
                pronunciation = extract_ipa(word_part)
                word_part = IPA_PATTERN.sub('', word_part).strip()
            # Extract part of speech
            pos_match = re.search(r'\(([nvadj]+p?)\)', word_part)
            if not pos_match:
                pos_match = re.search(r'\(([a-z]+)\)', word_part)
            if pos_match:
                candidate = pos_match.group(1)
                if candidate in ['n', 'v', 'a', 'adj', 'adv', 'np', 'phr']:
                    pos = candidate
                    word_part = word_part[:pos_match.start()].strip()
                elif len(candidate) <= 4:
                    pos = candidate
                    word_part = word_part[:pos_match.start()].strip()
            # Clean remaining word
            word = re.sub(r'\s+', ' ', word_part).strip()
            # Remove trailing special chars
            word = word.rstrip(' ,;:')
            break
        
        # Non-numbered but has IPA: "paint /peɪnt/" or "fishing /ˈfɪʃɪŋ/"
        if IPA_PATTERN.search(line) and not has_vietnamese(line):
            word_line_idx = i
            pronunciation = extract_ipa(line)
            word_part = IPA_PATTERN.sub('', line).strip()
            # Extract POS
            pos_match = re.search(r'\(([nvadj]+p?)\)', word_part)
            if pos_match:
                pos = pos_match.group(1)
                word_part = word_part[:pos_match.start()].strip()
            word = re.sub(r'\s+', ' ', word_part).strip().rstrip(' ,;:')
            break
    
    # If no word found yet, try the first non-Vietnamese line
    if not word and all_lines:
        for i, line in enumerate(all_lines):
            if not has_vietnamese(line) and not line.upper().startswith(tuple(kw.upper() for kw in TITLE_KEYWORDS)):
                word_line_idx = i
                if IPA_PATTERN.search(line):
                    pronunciation = extract_ipa(line)
                    line = IPA_PATTERN.sub('', line).strip()
                pos_match = re.search(r'\(([nvadj]+p?)\)', line)
                if pos_match:
                    pos = pos_match.group(1)
                    line = line[:pos_match.start()].strip()
                word = line.strip().rstrip(' ,;:')
                break
    
    if not word:
        return None
    
    # ===== PASS 2: Find IPA if not already found =====
    if not pronunciation:
        for i, line in enumerate(all_lines):
            if i == word_line_idx:
                continue
            ipa = extract_ipa(line)
            if ipa:
                pronunciation = ipa
                break
    
    # ===== PASS 3: Find meaning (Vietnamese text) =====
    for i, line in enumerate(all_lines):
        if i == word_line_idx:
            continue
        if has_vietnamese(line) and not line.lower().startswith('ex:'):
            # Check it's not an example translation in parentheses
            if not re.match(r'^\(.*\)$', line.strip()):
                meaning = line.strip()
                # Remove "Ex:" part if meaning and example are combined
                if 'Ex:' in meaning or 'ex:' in meaning:
                    parts = re.split(r'\n*[Ee]x:\s*', meaning, maxsplit=1)
                    meaning = parts[0].strip()
                    if len(parts) > 1:
                        example_en, example_vi = split_example(parts[1].strip())
                break
    
    # ===== PASS 4: Find example =====
    if not example_en:
        for i, line in enumerate(all_lines):
            if i == word_line_idx:
                continue
            # Explicit "Ex:" prefix
            if line.lower().startswith('ex:'):
                ex_text = line[3:].strip()
                example_en, example_vi = split_example(ex_text)
                break
            # English sentence (no Vietnamese, not the word line, not IPA, not meaning)
            if (not has_vietnamese(line) and i != word_line_idx 
                and not IPA_PATTERN.search(line) and line != meaning
                and len(line) > 10 and not line.startswith('(')):
                # Likely an example sentence
                example_en, example_vi = split_example(line)
                break
    
    # Also look for Vietnamese translation of example in remaining lines
    if example_en and not example_vi:
        for i, line in enumerate(all_lines):
            if i == word_line_idx:
                continue
            if has_vietnamese(line) and line != meaning and line.startswith('('):
                example_vi = line.strip('()')
                break
    
    # ===== Save best image =====
    image_path = None
    if images:
        best_img = max(images, key=lambda s: s.width * s.height if hasattr(s, 'width') else 0)
        image_path = save_image(best_img, lesson_id, word_index)
    
    # Clean up word - remove parenthetical usage notes like "(painting a house, a wall)"
    word = re.sub(r'\((?!n\b|v\b|adj\b|adv\b|np\b|a\b)[^)]*\)\s*$', '', word).strip()
    # Remove leading "To " if it's just an infinitive marker added to examples
    # But keep it if it's the actual word form
    
    return {
        'id': 'b%02d_w%02d' % (lesson_id, word_index),
        'word': word,
        'partOfSpeech': pos,
        'pronunciation': pronunciation,
        'meaning': meaning,
        'example': example_en,
        'exampleVi': example_vi,
        'image': image_path
    }


def process_pptx(filepath):
    filename = os.path.basename(filepath)
    lesson_id = extract_lesson_number(filename)
    
    print("\n  Processing: %s (Lesson %d)" % (filename, lesson_id))
    
    prs = Presentation(filepath)
    words = []
    word_index = 0
    
    for slide in prs.slides:
        word_index += 1
        result = parse_slide_unified(slide, lesson_id, word_index)
        if result:
            words.append(result)
        else:
            word_index -= 1
    
    print("    Extracted %d words" % len(words))
    
    # Show first word as sample
    if words:
        w = words[0]
        status = []
        if not w['pronunciation']:
            status.append('no-ipa')
        if not w['meaning']:
            status.append('no-meaning')
        status_str = ' [%s]' % ', '.join(status) if status else ''
        print("    Sample: %s %s %s%s" % (w['word'][:20], w['pronunciation'][:20], w['meaning'][:20], status_str))
    
    return {
        'id': lesson_id,
        'name': 'Bài %d' % lesson_id,
        'fileName': filename,
        'wordCount': len(words),
        'words': words
    }


def main():
    ensure_dirs()
    
    print("=" * 60)
    print("  VOCAB EXTRACTOR - PPTX to JSON + Images")
    print("=" * 60)
    
    pptx_files = sorted(
        [f for f in os.listdir(PPTX_DIR) if f.endswith('.pptx')],
        key=extract_lesson_number
    )
    
    print("\n  Found %d PPTX files in %s" % (len(pptx_files), PPTX_DIR))
    
    lessons = []
    total_words = 0
    
    for fname in pptx_files:
        filepath = os.path.join(PPTX_DIR, fname)
        try:
            lesson = process_pptx(filepath)
            if lesson['words']:
                lessons.append(lesson)
                total_words += len(lesson['words'])
        except Exception as e:
            print("    ERROR processing %s: %s" % (fname, e))
            import traceback
            traceback.print_exc()
    
    lessons.sort(key=lambda x: x['id'])
    
    output_data = {
        'totalLessons': len(lessons),
        'totalWords': total_words,
        'lessons': lessons
    }
    
    json_path = os.path.join(DATA_DIR, 'vocab.json')
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(output_data, f, ensure_ascii=False, indent=2)
    
    print("\n" + "=" * 60)
    print("  DONE!")
    print("  JSON: %s (%d words in %d lessons)" % (json_path, total_words, len(lessons)))
    print("  Images: %s (%d files)" % (IMAGES_DIR, len(os.listdir(IMAGES_DIR))))
    print("=" * 60)


if __name__ == '__main__':
    main()
